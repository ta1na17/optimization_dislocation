import base64
import logging
import re
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass, asdict
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reports.services.excel_layout import prepare_sheet_dataframe
from reports.services.owners import (
    OWNER_TO_TYPE,
    detect_owner_from_filename,
    detect_owner_from_text,
    normalize_canonical_owner,
)

logger = logging.getLogger('reports.builder')

BASE_STATIONS = {'ПСК': 'Калкаман', 'АлОр': 'Кызылорда', 'Индер': 'Макат'}
SHEET_ORDER = ['ПСК', 'АлОр', 'Индер']
DUPLICATE_FILL = PatternFill(start_color='FFF9C4', end_color='FFF9C4', fill_type='solid')

ITOG_STYLE_REFERENCE = Path(__file__).resolve().parent.parent / 'data' / 'itog_style_reference.xlsx'
ITOG_REF_SHEET_BY_OUTPUT = {'ПСК': 'ПСК1102', 'АлОр': 'АлОр1102', 'Индер': 'Индер1102'}

_THIN_BLACK = Side(style='thin', color='000000')
_FALLBACK_BORDER = Border(left=_THIN_BLACK, right=_THIN_BLACK, top=_THIN_BLACK, bottom=_THIN_BLACK)


def _open_itog_style_workbook():
    if not ITOG_STYLE_REFERENCE.exists():
        logger.warning('Эталон стилей итога не найден: %s', ITOG_STYLE_REFERENCE)
        return None
    try:
        return load_workbook(ITOG_STYLE_REFERENCE, data_only=False)
    except Exception as exc:
        logger.warning('Не удалось открыть эталон стилей: %s', exc)
        return None


def _copy_cell_style(dst, src):
    """Копирует оформление ячейки из эталона (openpyxl)."""
    if src is None:
        return
    if src.font:
        dst.font = copy(src.font)
    if src.border:
        dst.border = copy(src.border)
    if src.fill and getattr(src.fill, 'fill_type', None):
        dst.fill = copy(src.fill)
    if src.number_format:
        dst.number_format = src.number_format
    if src.protection:
        dst.protection = copy(src.protection)
    if src.alignment:
        dst.alignment = copy(src.alignment)


def _apply_column_widths_and_freeze(ws, ref_ws, ncols: int):
    if ref_ws.freeze_panes:
        ws.freeze_panes = ref_ws.freeze_panes
    else:
        ws.freeze_panes = 'A2'
    for i in range(1, ncols + 1):
        letter = get_column_letter(i)
        if letter in ref_ws.column_dimensions:
            w = ref_ws.column_dimensions[letter].width
            if w is not None:
                ws.column_dimensions[letter].width = w


def _apply_row_styles_from_reference(ws, ref_ws, ncols: int, segments: list[tuple[int, int, str]]):
    """
    segments: список (start_row, end_row, kind), kind = header | data | sep
    """
    ref_max_c = min(ref_ws.max_column, max(ncols, 11))
    for r1, r2, kind in segments:
        ref_row = 1 if kind == 'header' else 2
        for r in range(r1, r2 + 1):
            for c in range(1, ncols + 1):
                src_c = min(c, ref_max_c)
                _copy_cell_style(ws.cell(row=r, column=c), ref_ws.cell(row=ref_row, column=src_c))


@dataclass
class Row:
    wagon: str
    wagon_type: str | None
    payload: str | None
    operation_station: str | None
    destination_station: str | None
    distance_km: float | None
    idle_days: int | None
    operation: str | None
    owner: str | None
    source_file: str
    source_sheet: str | None
    source_row_number: int | None


COLUMN_ALIASES = {
    'wagon': [
        'номер вагона / контейнера',
        '№ вагона',
        'номер вагона',
        'вагон №',
        'вагон',
        'vagon',
        'wagon',
    ],
    'wagon_type': ['вид пс', 'тип вагона', 'тип модели', 'модель вагона'],
    'payload': [
        'г/п',
        'гп',
        'грузоподъёмность вагона',
        'грузоподъемность',
        'грузоподъемность (т.)',
        'грузоподъемность, тн',
        'вес груза, т.',
        'вес груза',
        'факт. вес',
        'вес',
    ],
    'operation_station': [
        'станция последней операции',
        'станция текущей дислокации',
        'нсп (дисл.)',
        'ст. операции',
        'станция опер.',
        'станция операции',
        'станция операции (рейс)',
    ],
    'destination_station': [
        'станция назначения',
        'ст. назначения',
        'ст. назн.',
        'назначение',
    ],
    'distance_km': [
        'расстояние до станции назначения',
        'текущ. расст. до ст. назн.',
        'остаточное расстояние',
        'ост. расстояние',
        'ост., км',
        'осталось, [км]',
        'осталось км',
        'расстояние осталось',
        'расстояние',
    ],
    'idle_days': [
        'дни без движения на ст. назначения',
        'дни без движения',
        'дней без движ.',
        'дней без движения',
        'простой с прибытия на станцию',
        'простой на станции дислокации',
        'простой от последней операции',
        'простой на станции',
        'простой',
        'без движения',
    ],
    'operation': ['операция', 'опер.', 'опер'],
    'owner': ['собственник', 'собств. вагона', 'владелец'],
    'capacity': [
        'фактический вес, тн',
        'фактический вес, т',
        'фактический вес',
        'вес груза, т.',
        'вес груза',
        'факт. вес',
        'объем,м3',
        'объем кузова',
        'объем вагона',
        'обьем вагона',
        'кубатура',
        'вместимость',
        'объем',
        'вес',
    ],
}


def _normalize_key(key: str) -> str:
    return ' '.join(str(key).strip().lower().split())


_UNIFY_SPACE = re.compile(r'[\u00a0\u1680\u2000-\u200b\u202f\u205f\u3000\ufeff]+')
_WAGON_MODEL_TYPE = re.compile(r'^\s*1([12])[-\s]')


def _normalize_station_label(value: str | None) -> str:
    """
    Станции из Excel: унифицируем «невидимые» пробелы и NBSP, схлопываем пробелы, регистр.
    В выгрузках dislocation станция часто с суффиксом дороги («Калкаман, КЗХ») — берём имя до запятой.
    Нужно для сопоставления с базовыми Калкаман / Кызылорда / Макат (ТЗ п.9–10.3).
    """
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ''
    s = str(value).replace('\ufeff', '')
    s = _UNIFY_SPACE.sub(' ', s)
    s = re.sub(r'\s+', ' ', s).strip().lower()
    if ',' in s:
        s = s.split(',', 1)[0].strip()
    return s


def _find_col(cols: list[str], aliases: list[str]) -> str | None:
    normalized = {_normalize_key(c): c for c in cols}
    for alias in aliases:
        if _normalize_key(alias) in normalized:
            return normalized[_normalize_key(alias)]
    # Подстрочное совпадение: выбираем столбец с самым длинным совпавшим псевдонимом (не «первый в файле»),
    # чтобы короткое «назначение» не перехватывало у «Ст. назначения» / «Код ст. назн.».
    by_len = sorted(aliases, key=lambda a: len(_normalize_key(a)), reverse=True)

    def longest_alias_in_name(nk: str) -> int:
        best = 0
        for alias in by_len:
            na = _normalize_key(alias)
            if len(na) < 2:
                continue
            if na in nk:
                best = max(best, len(na))
        return best

    best_col = None
    best_len = -1
    best_idx = 10**9
    for idx, c in enumerate(cols):
        nk = _normalize_key(c)
        L = longest_alias_in_name(nk)
        if L == 0:
            continue
        if L > best_len or (L == best_len and idx < best_idx):
            best_len = L
            best_idx = idx
            best_col = c
    return best_col


def _to_float(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, str):
        value = value.replace(',', '.').strip()
    try:
        return float(value)
    except Exception:
        return None


def _to_int(value):
    fv = _to_float(value)
    return None if fv is None else int(round(fv))


def _infer_type(owner, explicit_type, capacity):
    t, _ = _resolve_wagon_type(owner, explicit_type, capacity)
    return t


def _resolve_wagon_type(owner: str | None, explicit_type, capacity: float | None) -> tuple[str | None, list[str]]:
    """
    Тип вагона и список пояснений, если тип не выведен (для логов).
    Порядок: явный ПВ/КВ в таблице → OWNER_TO_TYPE → кубатура 120–161 (КВ), 70–90 (ПВ).
    """
    why: list[str] = []
    if explicit_type is not None and str(explicit_type).strip():
        raw = str(explicit_type).strip()
        t = raw.upper()
        if t in {'ПВ', 'КВ'}:
            return t, []
        model_m = _WAGON_MODEL_TYPE.match(raw)
        if model_m:
            return ('ПВ', 'КВ')[model_m.group(1) == '2'], []
        why.append(f'колонка типа/модели: значение {explicit_type!r} — не ПВ и не КВ')
    else:
        why.append('колонка типа/модели: пусто или столбец не сопоставлен с алиасами')
    if owner and owner in OWNER_TO_TYPE:
        return OWNER_TO_TYPE[owner], []
    if owner:
        why.append(f'собственник {owner!r} не входит в OWNER_TO_TYPE (маппинг ПВ/КВ только для части компаний)')
    else:
        why.append('собственник в строке не определён')
    if capacity is None:
        why.append('кубатура (capacity): столбец не найден или значение не приводится к числу')
    elif 120 <= capacity <= 161:
        return 'КВ', []
    elif 70 <= capacity <= 90:
        return 'ПВ', []
    else:
        why.append(f'кубатура (capacity) = {capacity} вне диапазонов КВ [120; 161] и ПВ [70; 90]')
    return None, why


def _row_snapshot_for_log(raw: pd.Series, col_names: list[str], max_len: int = 120) -> dict[str, str]:
    """Все ячейки строки по именам столбцов — для отладки пустого «Тип»."""
    out: dict[str, str] = {}
    for col in col_names:
        v = raw.get(col)
        if v is None or (isinstance(v, float) and pd.isna(v)):
            s = ''
        else:
            s = str(v).replace('\n', ' ').strip()
        if len(s) > max_len:
            s = s[: max_len - 3] + '...'
        out[str(col)] = s
    return out


def _parse_file(uploaded):
    ext = Path(uploaded.name).suffix.lower().strip()
    logger.info('Parsing file %s ext=%s size=%s', uploaded.name, ext, uploaded.size)
    content = uploaded.read()
    uploaded.seek(0)
    if ext == '.xlsx':
        return pd.read_excel(BytesIO(content), sheet_name=None, dtype=object, engine='openpyxl')
    if ext == '.xls':
        return pd.read_excel(BytesIO(content), sheet_name=None, dtype=object, engine='xlrd')
    raise ValueError(f'Формат файла не поддерживается: {uploaded.name}')


def _sheet_for_destination(destination: str | None) -> str | None:
    """
    Итоговый лист только если станция назначения совпадает с базовой станцией листа (ТЗ п.9).
    Иначе None — строка не включается в итог, а попадает в dropped_rows.
    """
    dest = _normalize_station_label(destination)
    if not dest:
        return None
    for sheet, station in BASE_STATIONS.items():
        if dest == _normalize_station_label(station):
            return sheet
    return None


DROP_REASON_NON_BASE_DESTINATION = (
    'станция назначения не совпадает ни с одной базовой станцией итоговых листов '
    '(Калкаман → лист ПСК, Кызылорда → лист АлОр, Макат → лист Индер) — строка в итог не включается'
)

# Секции в файлах с повтором шапки / «чужим» назначением: якорная строка с Калкаман/Кызылорда/Макат
# задаёт лист; строки с назначением = база идут в блоки 1–2 по ТЗ; с назначением ≠ базы — в блок 3 того же листа.
# Повтор шапки («Вагон №» / подпись столбца назначения) после секции АлОр: до строки МАКАТ небазовые строки — Индер блок 3.
DROP_REASON_SECTION_NO_ANCHOR = (
    'станция назначения не базовая и выше по файлу нет строки-якоря с Калкаман / Кызылорда / Макат '
    'для этой секции листа — строка в итог не включается'
)

_SUBTABLE_HEADER_WAGON = re.compile(r'(?i)^\s*вагон\s*№\s*$')


def _is_subtable_header_row(row: Row) -> bool:
    """Повторная шапка внутри листа (как в отчётах с двумя блоками вагонов)."""
    w = (row.wagon or '').strip()
    if w and _SUBTABLE_HEADER_WAGON.fullmatch(w):
        return True
    dnk = _normalize_key(row.destination_station or '')
    if not dnk:
        return False
    for alias in COLUMN_ALIASES['destination_station']:
        if dnk == _normalize_key(alias):
            return True
    return False


def _block_for_row(row: Row, sheet: str) -> int:
    op = _normalize_station_label(row.operation_station)
    dest = _normalize_station_label(row.destination_station)
    base = _normalize_station_label(BASE_STATIONS[sheet])
    all_base = {_normalize_station_label(v) for v in BASE_STATIONS.values()}
    if op and dest and op == dest:
        return 1
    if dest == base and op != dest:
        return 2
    if dest not in all_base:
        return 3
    return 3


def preview_owner_for_file(file_name: str, sheets: dict) -> dict:
    """
    Предпросмотр: какой собственник будет использован для файла до построчной сборки.
    Логика совпадает с приоритетом в _normalize_rows: сначала содержимое, затем имя файла.
    """
    for sheet_name, df in sheets.items():
        df = prepare_sheet_dataframe(df)
        if df is None or df.empty or df.shape[1] == 0:
            continue
        cols = [str(c) for c in df.columns]
        c_owner = _find_col(cols, COLUMN_ALIASES['owner'])
        max_rows = min(80, len(df))
        for idx in range(max_rows):
            raw = df.iloc[idx]
            row_num = int(idx) + 2
            if c_owner:
                o = detect_owner_from_text(raw.get(c_owner))
                if o:
                    logger.info('Preview owner for %s: %s (content, sheet=%s row=%s)', file_name, o, sheet_name, row_num)
                    return {'detected_owner': o, 'source': 'content', 'sheet_name': sheet_name, 'row_number': row_num}
            o = detect_owner_from_text(*[raw.get(col) for col in cols[: min(12, len(cols))]])
            if o:
                logger.info('Preview owner for %s: %s (content cells, sheet=%s row=%s)', file_name, o, sheet_name, row_num)
                return {'detected_owner': o, 'source': 'content', 'sheet_name': sheet_name, 'row_number': row_num}

    fn = detect_owner_from_filename(file_name)
    if fn:
        logger.info('Preview owner for %s: %s (filename)', file_name, fn)
        return {'detected_owner': fn, 'source': 'filename', 'sheet_name': None, 'row_number': None}

    logger.info('Preview owner for %s: not found', file_name)
    return {'detected_owner': None, 'source': 'unknown', 'sheet_name': None, 'row_number': None}


def preview_owners_for_uploads(files) -> list[dict]:
    out = []
    for f in files:
        try:
            sheets = _parse_file(f)
        except Exception as exc:
            logger.warning('Preview: cannot parse %s: %s', getattr(f, 'name', '?'), exc)
            out.append({
                'file_name': getattr(f, 'name', ''),
                'size': getattr(f, 'size', 0),
                'detected_owner': None,
                'source': 'error',
                'detail': str(exc),
            })
            continue
        info = preview_owner_for_file(f.name, sheets)
        out.append({
            'file_name': f.name,
            'size': f.size,
            **info,
        })
    return out


def validate_owner_choices_for_build(files, owner_overrides: list) -> str | None:
    """
    Если для файла собственник не определён автоматически (как в предпросмотре),
    в owner_overrides по индексу должно быть каноническое имя из справочника.
    """
    n = len(files)
    ov = list(owner_overrides) if owner_overrides else []
    if len(ov) < n:
        ov = ov + [None] * (n - len(ov))
    ov = ov[:n]

    for idx, f in enumerate(files):
        try:
            sheets = _parse_file(f)
        except Exception as exc:
            return f'Файл «{getattr(f, "name", "?")}» не удалось прочитать: {exc}'

        info = preview_owner_for_file(f.name, sheets)
        if info.get('source') != 'unknown':
            continue
        raw = ov[idx]
        if isinstance(raw, str) and not raw.strip():
            raw = None
        chosen = normalize_canonical_owner(raw) if raw else None
        if not chosen:
            return (
                f'Для файла «{getattr(f, "name", "?")}» собственник не определён автоматически. '
                'Выберите собственника из списка в столбце «Вручную».'
            )
    return None


def _normalize_rows(file_name: str, sheet_name: str, df: pd.DataFrame, file_owner_fallback: str | None = None):
    rows = []
    dropped = []
    df = prepare_sheet_dataframe(df)
    if df is None or df.empty or df.shape[1] == 0:
        return rows, dropped
    cols = [str(c) for c in df.columns]
    c = {k: _find_col(cols, v) for k, v in COLUMN_ALIASES.items()}
    binding_str = ', '.join(f'{k}→{c[k]!r}' for k in COLUMN_ALIASES)
    logger.info('Сопоставление столбцов %s / %s: %s', file_name, sheet_name, binding_str)

    for idx, raw in df.iterrows():
        row_num = int(idx) + 2
        wagon = raw.get(c['wagon']) if c['wagon'] else None
        wagon_text = '' if wagon is None or (isinstance(wagon, float) and pd.isna(wagon)) else str(wagon).strip()
        if not wagon_text:
            continue
        wt_lower = wagon_text.lower()
        if wt_lower in {'итого', 'всего', 'nan'} or 'итого' in wt_lower:
            continue
        digits_only = ''.join(ch for ch in wagon_text if ch.isdigit())
        if digits_only and len(digits_only) < 6:
            continue

        station_op = raw.get(c['operation_station']) if c['operation_station'] else None
        station_dest = raw.get(c['destination_station']) if c['destination_station'] else None
        operation = raw.get(c['operation']) if c['operation'] else None
        owner = None
        if c['owner']:
            owner = detect_owner_from_text(raw.get(c['owner']))
        if not owner:
            owner = detect_owner_from_text(*[raw.get(col) for col in cols[:min(12, len(cols))]])
        if not owner and file_owner_fallback:
            owner = file_owner_fallback
        if not owner:
            owner = detect_owner_from_filename(file_name)

        minimal = [station_op, station_dest, operation, raw.get(c['distance_km']) if c['distance_km'] else None, raw.get(c['idle_days']) if c['idle_days'] else None]
        non_empty = [v for v in minimal if v is not None and not (isinstance(v, float) and pd.isna(v)) and str(v).strip()]
        if not non_empty:
            dropped.append({
                'file_name': file_name,
                'sheet_name': sheet_name,
                'row_number': row_num,
                'row_preview': f'Вагон={wagon_text}',
                'reason': 'строка содержит только номер вагона',
            })
            continue

        if not station_dest or (isinstance(station_dest, float) and pd.isna(station_dest)):
            dropped.append({
                'file_name': file_name,
                'sheet_name': sheet_name,
                'row_number': row_num,
                'row_preview': f'Вагон={wagon_text}',
                'reason': 'не удалось определить обязательные поля',
            })
            continue

        payload_value = raw.get(c['payload']) if c['payload'] else None
        cap_raw = raw.get(c['capacity']) if c['capacity'] else None
        capacity_float = _to_float(cap_raw)
        wagon_type_raw = raw.get(c['wagon_type']) if c['wagon_type'] else None
        inferred_type, type_why = _resolve_wagon_type(owner, wagon_type_raw, capacity_float)
        payload_float = _to_float(payload_value)

        if inferred_type is None:
            snap = _row_snapshot_for_log(raw, cols)
            owner_col_raw = raw.get(c['owner']) if c['owner'] else None
            logger.warning(
                'ТИП_НЕ_ОПРЕДЕЛЁН file=%s sheet=%s row_excel≈%s wagon=%r resolved_owner=%r | причины: %s | '
                'колонки: owner=%r raw=%r | wagon_type=%r raw=%r | capacity=%r raw=%r float=%s | '
                'payload=%r raw=%r float=%s | полная_строка=%s',
                file_name,
                sheet_name,
                row_num,
                wagon_text,
                owner,
                ' | '.join(type_why),
                c['owner'],
                owner_col_raw,
                c['wagon_type'],
                wagon_type_raw,
                c['capacity'],
                cap_raw,
                capacity_float,
                c['payload'],
                payload_value,
                payload_float,
                snap,
            )
            if payload_float is not None and capacity_float is None:
                logger.warning(
                    'ТИП_ПОДСКАЗКА: в г/п есть число %s, но для «Тип» по ТЗ используется кубатура (capacity), не вес груза; '
                    'проверьте сопоставление столбца кубатуры file=%s row≈%s',
                    payload_float,
                    file_name,
                    row_num,
                )

        rows.append(Row(
            wagon=wagon_text,
            wagon_type=inferred_type,
            payload='' if payload_value is None or (isinstance(payload_value, float) and pd.isna(payload_value)) else str(payload_value),
            operation_station='' if station_op is None or (isinstance(station_op, float) and pd.isna(station_op)) else str(station_op).strip(),
            destination_station=str(station_dest).strip(),
            distance_km=_to_float(raw.get(c['distance_km'])) if c['distance_km'] else None,
            idle_days=_to_int(raw.get(c['idle_days'])) if c['idle_days'] else None,
            operation='' if operation is None or (isinstance(operation, float) and pd.isna(operation)) else str(operation).strip(),
            owner=owner,
            source_file=file_name,
            source_sheet=sheet_name,
            source_row_number=row_num,
        ))

    logger.info('Normalized sheet %s/%s: rows=%s dropped=%s', file_name, sheet_name, len(rows), len(dropped))
    return rows, dropped


def _fallback_grid_style(ws, ncols: int, header_rows: set[int], max_row: int):
    """Если нет эталона — тонкая сетка и жирная шапка."""
    ws.freeze_panes = 'A2'
    for r in range(1, max_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = _FALLBACK_BORDER
            if r in header_rows:
                cell.font = Font(bold=True)


def _write_excel(grouped: dict, out: Path | BytesIO) -> int:
    wb = Workbook()
    wb.remove(wb.active)
    all_rows = [r for sheet in SHEET_ORDER for block in (1, 2, 3) for r in grouped[sheet][block]]
    counts = Counter(r.wagon for r in all_rows if r.wagon)
    duplicates = {w for w, cnt in counts.items() if cnt > 1}

    headers_common = ['Вагон', 'Тип', 'Станция операции', 'Станция назначения', 'Осталось, [км]', 'Дни без движения', 'Операция', 'Собственник', 'Примечание', 'Статус']
    headers_psk = ['Вагон', 'Тип', 'г/п', 'Станция операции', 'Станция назначения', 'Осталось, [км]', 'Дни без движения', 'Операция', 'Собственник', 'Примечание', 'Статус']

    ref_wb = _open_itog_style_workbook()

    for sheet in SHEET_ORDER:
        ws = wb.create_sheet(sheet)
        headers = headers_psk if sheet == 'ПСК' else headers_common
        ncols = len(headers)
        segments: list[tuple[int, int, str]] = []
        header_rows: set[int] = set()
        row_idx = 1

        for block in (1, 2, 3):
            segments.append((row_idx, row_idx, 'header'))
            header_rows.add(row_idx)
            for i, h in enumerate(headers, start=1):
                ws.cell(row=row_idx, column=i, value=h)
            row_idx += 1

            block_rows = grouped[sheet][block]
            if block_rows:
                data_start = row_idx
                for r in block_rows:
                    values = [r.wagon, r.wagon_type]
                    if sheet == 'ПСК':
                        values.append(r.payload)
                    values.extend([r.operation_station, r.destination_station, r.distance_km, r.idle_days, r.operation, r.owner, '', ''])
                    for i, v in enumerate(values, start=1):
                        ws.cell(row=row_idx, column=i, value=v)
                    row_idx += 1
                segments.append((data_start, row_idx - 1, 'data'))
            if block < 3:
                sep_row = row_idx
                for i in range(1, ncols + 1):
                    ws.cell(row=sep_row, column=i, value='X' if i == 1 else None)
                segments.append((sep_row, sep_row, 'sep'))
                row_idx += 1

        ref_ws = None
        if ref_wb:
            ref_name = ITOG_REF_SHEET_BY_OUTPUT.get(sheet)
            if ref_name and ref_name in ref_wb.sheetnames:
                ref_ws = ref_wb[ref_name]

        if ref_ws:
            _apply_column_widths_and_freeze(ws, ref_ws, ncols)
            _apply_row_styles_from_reference(ws, ref_ws, ncols, segments)
        else:
            _fallback_grid_style(ws, ncols, header_rows, row_idx - 1)

        for r in range(1, row_idx):
            v0 = ws.cell(row=r, column=1).value
            if v0 is None or str(v0).strip().upper() == 'X':
                continue
            if v0 in duplicates:
                for c in range(1, ncols + 1):
                    ws.cell(row=r, column=c).fill = DUPLICATE_FILL

    wb.save(out)
    logger.info('Excel written: %s duplicates=%s', out, len(duplicates))
    return len(duplicates)


def build_report(files, owner_overrides: list[str | None] | None = None):
    """
    owner_overrides — по порядку с файлами: каноническое имя собственника или None.
    Подставляется в строку, если в Excel не найден собственник: сначала значение «Вручную»,
    иначе то же автоопределение, что в предпросмотре (по содержимому или имени файла).
    """
    logger.info('Build started, files=%s overrides=%s', len(files), owner_overrides)
    all_rows = []
    dropped = []

    if owner_overrides is None:
        owner_overrides = [None] * len(files)
    elif len(owner_overrides) != len(files):
        logger.warning('owner_overrides length %s != files %s, padding', len(owner_overrides), len(files))
        owner_overrides = list(owner_overrides) + [None] * max(0, len(files) - len(owner_overrides))
        owner_overrides = owner_overrides[: len(files)]

    for idx, f in enumerate(files):
        manual = normalize_canonical_owner(owner_overrides[idx]) if owner_overrides[idx] else None
        try:
            sheets = _parse_file(f)
        except Exception as exc:
            logger.warning('Failed parsing file %s: %s', getattr(f, 'name', 'unknown'), exc)
            return None, str(exc)

        preview_info = preview_owner_for_file(f.name, sheets)
        auto_owner = None
        if preview_info.get('source') in ('content', 'filename') and preview_info.get('detected_owner'):
            auto_owner = preview_info['detected_owner']
        file_owner_fallback = manual or auto_owner

        for sheet_name, df in sheets.items():
            rows, drop_rows = _normalize_rows(f.name, sheet_name, df, file_owner_fallback=file_owner_fallback)
            all_rows.extend(rows)
            dropped.extend(drop_rows)

    grouped = {s: defaultdict(list) for s in SHEET_ORDER}
    rows_in_output = 0
    # (файл, лист Excel) → (нормализованное назначение якоря, итоговый лист ПСК/АлОр/Индер)
    section_anchor: dict[tuple[str, str], tuple[str, str]] = {}
    # После внутренней шапки в секции АлОр: небазовые строки до первой строки МАКАТ — в Индер блок 3
    pre_makat_tail: dict[tuple[str, str], bool] = {}

    for row in all_rows:
        key = (row.source_file, row.source_sheet)
        if _is_subtable_header_row(row):
            anchor = section_anchor.get(key)
            if anchor and anchor[1] == 'АлОр':
                pre_makat_tail[key] = True
            # Служебная строка: не в итог и не в списке «отброшенных» (не ошибка данных)
            continue

        dest_norm = _normalize_station_label(row.destination_station)
        sheet_from_dest = _sheet_for_destination(row.destination_station)

        if sheet_from_dest is not None:
            pre_makat_tail[key] = False
            section_anchor[key] = (dest_norm, sheet_from_dest)
            block = _block_for_row(row, sheet_from_dest)
            grouped[sheet_from_dest][block].append(row)
            rows_in_output += 1
            continue

        if pre_makat_tail.get(key):
            grouped['Индер'][3].append(row)
            rows_in_output += 1
            continue

        anchor = section_anchor.get(key)
        if anchor is not None:
            anchor_dest, anchor_sheet = anchor
            if dest_norm != anchor_dest:
                grouped[anchor_sheet][3].append(row)
                rows_in_output += 1
                continue

        dropped.append(
            {
                'file_name': row.source_file,
                'sheet_name': row.source_sheet,
                'row_number': row.source_row_number,
                'row_preview': (
                    f'Вагон={row.wagon}, ст.опер.={row.operation_station!r}, ст.назн.={row.destination_station!r}'
                ),
                'reason': (
                    DROP_REASON_SECTION_NO_ANCHOR
                    if not dest_norm
                    else DROP_REASON_NON_BASE_DESTINATION
                ),
            }
        )

    for sheet in SHEET_ORDER:
        grouped[sheet][1].sort(key=lambda r: (r.idle_days is None, -(r.idle_days or 0)))
        grouped[sheet][2].sort(key=lambda r: (r.distance_km is None, r.distance_km or 0))
        grouped[sheet][3].sort(key=lambda r: (r.idle_days is None, -(r.idle_days or 0)))

    buf = BytesIO()
    duplicates_count = _write_excel(grouped, buf)
    file_bytes = buf.getvalue()

    result = {
        'duplicates_count': duplicates_count,
        'dropped_rows_count': len(dropped),
        'dropped_rows': dropped,
        'file_base64': base64.b64encode(file_bytes).decode('ascii'),
        'file_name': 'Итог.xlsx',
    }
    logger.info(
        'Build complete: normalized=%s in_output=%s dropped=%s duplicates=%s bytes=%s',
        len(all_rows),
        rows_in_output,
        len(dropped),
        duplicates_count,
        len(file_bytes),
    )
    return result, None
