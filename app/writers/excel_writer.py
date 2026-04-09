from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.rules.distribution import SHEET_ORDER


DUPLICATE_FILL = PatternFill(start_color='FFF9C4', end_color='FFF9C4', fill_type='solid')
HEADERS_COMMON = ['Вагон', 'Тип', 'Станция операции', 'Станция назначения', 'Осталось, [км]', 'Дни без движения', 'Операция', 'Собственник', 'Примечание', 'Статус']
HEADERS_PSK = ['Вагон', 'Тип', 'г/п', 'Станция операции', 'Станция назначения', 'Осталось, [км]', 'Дни без движения', 'Операция', 'Собственник', 'Примечание', 'Статус']


def write_result(grouped: dict, output_path: Path) -> int:
    wb = Workbook()
    wb.remove(wb.active)

    all_rows = []
    for sheet in SHEET_ORDER:
        for block in (1, 2, 3):
            all_rows.extend(grouped[sheet][block])
    wagon_counts = Counter(r.wagon for r in all_rows if r.wagon)
    duplicates = {w for w, c in wagon_counts.items() if c > 1}

    for sheet in SHEET_ORDER:
        ws = wb.create_sheet(sheet)
        headers = HEADERS_PSK if sheet == 'ПСК' else HEADERS_COMMON

        current_row = 1
        for block in (1, 2, 3):
            for col, h in enumerate(headers, start=1):
                cell = ws.cell(row=current_row, column=col, value=h)
                cell.font = Font(bold=True)
            current_row += 1

            for r in grouped[sheet][block]:
                values = [
                    r.wagon,
                    r.wagon_type,
                ]
                if sheet == 'ПСК':
                    values.append(r.payload)
                values.extend([
                    r.operation_station,
                    r.destination_station,
                    r.distance_km,
                    r.idle_days,
                    r.operation,
                    r.owner,
                    '',
                    '',
                ])

                for col, val in enumerate(values, start=1):
                    ws.cell(row=current_row, column=col, value=val)

                if r.wagon in duplicates:
                    for col in range(1, len(headers) + 1):
                        ws.cell(row=current_row, column=col).fill = DUPLICATE_FILL

                current_row += 1

            if block < 3:
                ws.cell(row=current_row, column=1, value='X')
                current_row += 1

        widths = [15, 10, 12, 24, 24, 14, 18, 18, 22, 16, 12]
        for i, w in enumerate(widths, start=1):
            if i <= len(headers):
                ws.column_dimensions[chr(64 + i)].width = w

    wb.save(output_path)
    return len(duplicates)
